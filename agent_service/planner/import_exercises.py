# -*- coding: utf-8 -*-
"""
Import exercises from workout_details.xlsx → PostgreSQL exercises table.

Pipeline:
  1. Read unique exercises from xlsx
  2. Map deterministic fields (muscles, difficulty, type, equipment …)
  3. Call Gemini API in batches of 5 to generate:
       name_cn, secondary_muscles, common_mistakes, safety_tips,
       contraindications, fatigue_score
  4. Upsert into exercises table
"""

import json
import os
import re
import sys
import time
import traceback
from pathlib import Path

import openpyxl
import psycopg2
import psycopg2.extras
from openai import OpenAI

# ──────────────────────────────────────────────
# Config
# ──────────────────────────────────────────────
BASE_DIR   = Path("D:/WorkPath/fitness/agent_service/planner")
XLSX_PATH  = BASE_DIR / "workout_details.xlsx"
SQL_PATH   = BASE_DIR / "create_exercises_table.sql"

DB_CFG = dict(host="localhost", port=5432,
              user="postgres", password="666666", dbname="fitness")

GEMINI_BASE_URL = "https://nextrouter.cc/v1"
GEMINI_API_KEY  = "sk-gxBoLiZEqsQnwPweg2mlV6AWz0gpQJzlbtFz2rzRBYova4Fc"
GEMINI_MODEL    = "gemini-3.5-flash"
BATCH_SIZE      = 5   # exercises per LLM call

# ──────────────────────────────────────────────
# Lookup tables
# ──────────────────────────────────────────────
MUSCLE_ZH = {
    "quads":           "股四头肌",
    "glutes":          "臀大肌",
    "hamstrings":      "腘绳肌",
    "calves":          "腓肠肌",
    "chest":           "胸大肌",
    "triceps":         "肱三头肌",
    "biceps":          "肱二头肌",
    "lats":            "背阔肌",
    "traps":           "斜方肌",
    "traps-middle":    "斜方肌中部",
    "abdominals":      "腹直肌",
    "obliques":        "腹斜肌",
    "forearms":        "前臂肌群",
    "hands":           "手部肌群",
    "front-shoulders": "三角肌前束",
    "rear-shoulders":  "三角肌后束",
    "shoulders":       "三角肌中束",
    "lowerback":       "竖脊肌",
    "shins":           "胫骨前肌",
    "neck":            "颈部肌群",
}

DIFFICULTY_MAP = {
    "Novice":       "beginner",
    "Beginner":     "beginner",
    "Intermediate": "intermediate",
    "Advanced":     "advanced",
}

MOVEMENT_MAP = {
    "Push": "push",
    "Pull": "pull",
    "Hold": "hold",
}

TYPE_MAP = {
    "Compound":  "compound",
    "Isolation": "isolation",
}

EQUIPMENT_KEYWORDS = [
    ("dumbbell",       "哑铃"),
    ("barbell",        "杠铃"),
    ("ez bar",         "EZ杠"),
    ("trap bar",       "六角杠"),
    ("cable",          "缆绳机"),
    ("smith",          "史密斯架"),
    ("kettlebell",     "壶铃"),
    ("band",           "弹力带"),
    ("machine",        "器械"),
    ("pull up",        "单杠"),
    ("pull-up",        "单杠"),
    ("pullup",         "单杠"),
    ("chin up",        "单杠"),
    ("chin-up",        "单杠"),
    ("dip",            "双杠"),
    ("suspension",     "悬吊绳"),
    ("bosu",           "波速球"),
    ("stability ball", "瑞士球"),
    ("foam",           "泡沫轴"),
]
BODYWEIGHT_KEYWORDS = [
    "plank", "push up", "push-up", "pushup", "burpee",
    "crunch", "sit up", "sit-up", "leg raise", "mountain climber",
    "jumping jack", "squat jump", "glute bridge", "hip bridge",
    "superman", "bird dog", "dead bug",
]

# Needs bench
BENCH_KEYWORDS = ["bench", "incline", "decline", "flat", " fly", "flye", "press"]


def detect_equipment(name: str) -> list:
    nl = name.lower()
    equip = []
    for kw, zh in EQUIPMENT_KEYWORDS:
        if kw in nl and zh not in equip:
            equip.append(zh)
    is_bw = any(kw in nl for kw in BODYWEIGHT_KEYWORDS)
    if not equip or is_bw:
        equip.insert(0, "自重")
    if any(kw in nl for kw in BENCH_KEYWORDS) and "杠铃" not in equip and "哑铃" not in equip:
        equip.append("训练凳")
    elif "杠铃" in equip or "哑铃" in equip:
        if any(kw in nl for kw in ["bench", "incline", "decline", "flat"]):
            equip.append("训练凳")
    return list(dict.fromkeys(equip))  # preserve order, deduplicate


def derive_risk(name: str, difficulty: str, equip: list, ex_type: str) -> str:
    if difficulty == "advanced" and "杠铃" in equip:
        return "high"
    if "杠铃" in equip or ex_type == "compound":
        return "medium"
    if "器械" in equip:
        return "low"
    if difficulty == "advanced":
        return "medium"
    return "low"


def derive_mets(ex_type: str, difficulty: str, equip: list) -> float:
    base = 5.0 if ex_type == "compound" else 3.5
    bonus = {"beginner": 0, "intermediate": 0.5, "advanced": 1.5}.get(difficulty, 0)
    if "杠铃" in equip:
        bonus += 1.0
    return round(base + bonus, 1)


def derive_training_goals(movement: str, ex_type: str, difficulty: str) -> list:
    if movement == "hold":
        return ["核心稳定", "耐力提升"]
    if ex_type == "compound":
        if difficulty == "advanced":
            return ["力量提升", "增肌"]
        return ["增肌", "力量提升"]
    return ["增肌", "塑形"]


def make_exercise_id(name: str, seen_ids: set) -> str:
    base = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")[:60]
    cand = f"{base}_001"
    idx  = 1
    while cand in seen_ids:
        idx += 1
        cand = f"{base}_{idx:03d}"
    seen_ids.add(cand)
    return cand


def get_primary_muscles(muscles_dict: dict) -> list:
    """Collect text-mw-red muscle IDs → Chinese names."""
    red = []
    for side in ["frontBodyMap", "backBodyMap"]:
        for mid in muscles_dict.get(side, {}).get("text-mw-red", []):
            zh = MUSCLE_ZH.get(mid)
            if zh and zh not in red:
                red.append(zh)
    return red


# ──────────────────────────────────────────────
# Read & deduplicate exercises
# ──────────────────────────────────────────────
def load_unique_exercises(xlsx_path: Path) -> list:
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb.active
    seen_names = {}
    for r in range(2, ws.max_row + 1):
        name    = (ws.cell(r, 2).value or "").strip()
        details = ws.cell(r, 4).value or "{}"
        if name and name not in seen_names:
            seen_names[name] = json.loads(details)
    result = [{"name": n, "details": d} for n, d in seen_names.items()]
    print(f"Unique exercises loaded: {len(result)}")
    return result


# ──────────────────────────────────────────────
# LLM enrichment
# ──────────────────────────────────────────────
LLM_SYSTEM = """你是专业健身教练兼运动科学专家。
用户会提供若干健身动作，你需要为每个动作返回结构化 JSON 数据（中文），
确保内容专业、实用、简洁。"""

LLM_USER_TMPL = """\
请为以下 {n} 个动作逐一生成 JSON，以 JSON 数组格式返回（顺序与输入一致）。

每个对象的格式：
{{
  "name_cn": "动作中文名",
  "secondary_muscles": ["次要肌群1", "次要肌群2"],
  "common_mistakes": ["错误1", "错误2", "错误3"],
  "safety_tips": ["提示1", "提示2"],
  "contraindications": ["禁忌1", "禁忌2"],
  "estimated_mets": 6.0,
  "fatigue_score": {{"目标肌群": 8, "中枢神经": 5}}
}}

注意：
- fatigue_score 的 key 用中文肌群名，值 1-10
- 中枢神经疲劳：复合动作 6-8，孤立动作 2-4
- estimated_mets 范围 2-12
- 禁忌 / 常见错误 各 2-4 条，安全提示 2-3 条

动作列表：
{exercises_block}

只返回 JSON 数组，不要添加任何其他文字。"""


def call_llm(client: OpenAI, batch: list) -> list:
    """Call Gemini for a batch; return list of enrichment dicts.

    Validates response data quality and retries on dirty data (non-string
    values in lists, illegal characters, concatenated muscle IDs, etc.).
    """
    lines = []
    for i, ex in enumerate(batch, 1):
        pm = ", ".join(ex["primary_muscles"]) or "未知"
        lines.append(
            f"{i}. 动作名={ex['name']}  主要肌群={pm}  "
            f"动作模式={ex.get('movement_pattern','未知')}  "
            f"难度={ex.get('difficulty','beginner')}  "
            f"类型={ex.get('exercise_type','compound')}"
        )
    block  = "\n".join(lines)
    prompt = LLM_USER_TMPL.format(n=len(batch), exercises_block=block)

    MAX_ATTEMPTS = 5
    for attempt in range(MAX_ATTEMPTS):
        try:
            resp = client.chat.completions.create(
                model=GEMINI_MODEL,
                messages=[
                    {"role": "system", "content": LLM_SYSTEM},
                    {"role": "user",   "content": prompt},
                ],
                temperature=0.3,
                timeout=60,
            )
            raw = resp.choices[0].message.content.strip()
            raw = re.sub(r"^```(?:json)?\s*", "", raw)
            raw = re.sub(r"\s*```$", "", raw)
            data = json.loads(raw)
            if not isinstance(data, list) or len(data) != len(batch):
                print(f"  [LLM] Length mismatch: got {len(data) if isinstance(data, list) else type(data).__name__}, expected {len(batch)}")
                time.sleep(2)
                continue

            # --- Validate each item ---
            all_errors: list[str] = []
            for idx, item in enumerate(data):
                if not isinstance(item, dict):
                    all_errors.append(f"  item[{idx}] is not a dict")
                    continue
                errs = _validate_enrichment(item)
                for e in errs:
                    all_errors.append(f"  item[{idx}] ({batch[idx]['name']}): {e}")

            if all_errors and attempt < MAX_ATTEMPTS - 1:
                print(f"  [LLM] Attempt {attempt+1}: dirty data detected ({len(all_errors)} issues):")
                for e in all_errors[:5]:
                    print(f"    {e}")
                if len(all_errors) > 5:
                    print(f"    ... and {len(all_errors) - 5} more")
                time.sleep(2)
                continue

            if all_errors:
                print(f"  [LLM] Last attempt still has {len(all_errors)} issues, sanitizing in-place")

            # --- Sanitize: fix recoverable dirty data ---
            for item in data:
                if not isinstance(item, dict):
                    continue
                if "secondary_muscles" in item:
                    item["secondary_muscles"] = _sanitize_muscle_list(
                        item["secondary_muscles"] if isinstance(item["secondary_muscles"], list) else []
                    )
                for field in ("common_mistakes", "safety_tips", "contraindications"):
                    val = item.get(field)
                    if isinstance(val, list):
                        item[field] = [
                            _sanitize_str(str(v)) for v in val
                            if isinstance(v, str) and _sanitize_str(v)
                        ]

            return data
        except json.JSONDecodeError as e:
            print(f"  [LLM] Attempt {attempt+1} JSON parse failed: {e}")
            time.sleep(2)
        except Exception as e:
            print(f"  [LLM] Attempt {attempt+1} failed: {e}")
            time.sleep(2)
    # Fallback: empty enrichment for each
    return [{}] * len(batch)


VALID_MUSCLE_NAMES_ZH = set(MUSCLE_ZH.values())

VALID_MUSCLE_IDS = set(MUSCLE_ZH.keys())

_MUSCLE_CLEAN_RE = re.compile(r'[?\n\r\t\x00-\x1f]')


def _sanitize_str(v: str) -> str:
    """Strip illegal characters from a string value."""
    return _MUSCLE_CLEAN_RE.sub('', v).strip()


def _sanitize_muscle_list(raw: list) -> list[str]:
    """Clean a muscle list: ensure all items are strings, fix common Gemini errors."""
    cleaned = []
    for item in raw:
        if not isinstance(item, str):
            continue
        s = _sanitize_str(item)
        if '_' in s and s not in VALID_MUSCLE_IDS:
            parts = s.split('_')
            for p in parts:
                p = p.strip()
                if p and (p in VALID_MUSCLE_NAMES_ZH or p in VALID_MUSCLE_IDS):
                    cleaned.append(p)
            continue
        if s:
            cleaned.append(s)
    return list(dict.fromkeys(cleaned))


def _validate_enrichment(item: dict) -> list[str]:
    """Validate a single Gemini enrichment result. Returns a list of error messages."""
    errors = []
    sm = item.get("secondary_muscles", [])
    if not isinstance(sm, list):
        errors.append(f"secondary_muscles is not a list: {type(sm).__name__}")
    else:
        for i, v in enumerate(sm):
            if not isinstance(v, str):
                errors.append(f"secondary_muscles[{i}] is {type(v).__name__}({v!r}), expected str")
            elif _MUSCLE_CLEAN_RE.search(v):
                errors.append(f"secondary_muscles[{i}] contains illegal chars: {v!r}")
            elif '_' in v and v not in VALID_MUSCLE_IDS:
                errors.append(f"secondary_muscles[{i}] looks like concatenated IDs: {v!r}")

    cm = item.get("common_mistakes", [])
    if not isinstance(cm, list):
        errors.append(f"common_mistakes is not a list: {type(cm).__name__}")

    st = item.get("safety_tips", [])
    if not isinstance(st, list):
        errors.append(f"safety_tips is not a list: {type(st).__name__}")

    mets = item.get("estimated_mets")
    if mets is not None:
        try:
            fv = float(mets)
            if not (1.0 <= fv <= 15.0):
                errors.append(f"estimated_mets out of range: {fv}")
        except (TypeError, ValueError):
            errors.append(f"estimated_mets not numeric: {mets!r}")

    return errors


def safe_list(v) -> list:
    if isinstance(v, list):
        return v
    return []

def safe_float(v, default=5.0) -> float:
    try:
        return float(v)
    except Exception:
        return default

def safe_dict(v) -> dict:
    if isinstance(v, dict):
        return v
    return {}


# ──────────────────────────────────────────────
# DB helpers
# ──────────────────────────────────────────────
UPSERT_SQL = """
INSERT INTO exercises (
    exercise_id, name, name_cn,
    primary_muscles, secondary_muscles,
    movement_pattern, equipment, difficulty,
    training_goals, exercise_type,
    recommended_rep_range, risk_level,
    common_mistakes, safety_tips,
    estimated_mets, fatigue_score,
    contraindications, source_data
) VALUES (
    %(exercise_id)s, %(name)s, %(name_cn)s,
    %(primary_muscles)s::jsonb, %(secondary_muscles)s::jsonb,
    %(movement_pattern)s, %(equipment)s::jsonb, %(difficulty)s,
    %(training_goals)s::jsonb, %(exercise_type)s,
    %(recommended_rep_range)s::jsonb, %(risk_level)s,
    %(common_mistakes)s::jsonb, %(safety_tips)s::jsonb,
    %(estimated_mets)s, %(fatigue_score)s::jsonb,
    %(contraindications)s::jsonb, %(source_data)s::jsonb
)
ON CONFLICT (exercise_id) DO UPDATE SET
    name_cn              = EXCLUDED.name_cn,
    primary_muscles      = EXCLUDED.primary_muscles,
    secondary_muscles    = EXCLUDED.secondary_muscles,
    movement_pattern     = EXCLUDED.movement_pattern,
    equipment            = EXCLUDED.equipment,
    difficulty           = EXCLUDED.difficulty,
    training_goals       = EXCLUDED.training_goals,
    exercise_type        = EXCLUDED.exercise_type,
    recommended_rep_range= EXCLUDED.recommended_rep_range,
    risk_level           = EXCLUDED.risk_level,
    common_mistakes      = EXCLUDED.common_mistakes,
    safety_tips          = EXCLUDED.safety_tips,
    estimated_mets       = EXCLUDED.estimated_mets,
    fatigue_score        = EXCLUDED.fatigue_score,
    contraindications    = EXCLUDED.contraindications,
    source_data          = EXCLUDED.source_data,
    updated_at           = NOW()
"""


def run_ddl(conn):
    sql = SQL_PATH.read_text(encoding="utf-8")
    with conn.cursor() as cur:
        cur.execute(sql)
    conn.commit()
    print("Table & indexes created/verified.")


# ──────────────────────────────────────────────
# Main pipeline
# ──────────────────────────────────────────────
def main():
    # 1. DB connection & DDL
    print("Connecting to PostgreSQL …")
    conn = psycopg2.connect(**DB_CFG)
    run_ddl(conn)

    # 2. LLM client
    llm = OpenAI(base_url=GEMINI_BASE_URL, api_key=GEMINI_API_KEY)

    # 3. Load exercises
    raw_exercises = load_unique_exercises(XLSX_PATH)

    # 4. Build deterministic fields
    seen_ids = set()
    exercises = []
    for ex in raw_exercises:
        bi   = ex["details"].get("basic_info", {})
        mmap = ex["details"].get("muscles", {})

        # Find field values by matching known English values
        difficulty_raw = movement_raw = ex_type_raw = ""
        for v in bi.values():
            if v in DIFFICULTY_MAP:  difficulty_raw = v
            elif v in MOVEMENT_MAP:  movement_raw   = v
            elif v in TYPE_MAP:      ex_type_raw    = v

        difficulty  = DIFFICULTY_MAP.get(difficulty_raw, "beginner")
        movement    = MOVEMENT_MAP.get(movement_raw,    "push")
        ex_type     = TYPE_MAP.get(ex_type_raw,         "compound")
        primary_zh  = get_primary_muscles(mmap)
        equip       = detect_equipment(ex["name"])
        risk        = derive_risk(ex["name"], difficulty, equip, ex_type)
        mets        = derive_mets(ex_type, difficulty, equip)
        goals       = derive_training_goals(movement, ex_type, difficulty)
        ex_id       = make_exercise_id(ex["name"], seen_ids)

        exercises.append({
            "exercise_id":    ex_id,
            "name":           ex["name"],
            "primary_muscles": primary_zh,
            "movement_pattern": movement,
            "equipment":       equip,
            "difficulty":      difficulty,
            "training_goals":  goals,
            "exercise_type":   ex_type,
            "risk_level":      risk,
            "estimated_mets":  mets,
            "recommended_rep_range": {
                "strength":    "3-6",
                "hypertrophy": "8-12",
                "endurance":   "15-20",
            },
            # LLM fields — filled next
            "name_cn":           "",
            "secondary_muscles": [],
            "common_mistakes":   [],
            "safety_tips":       [],
            "contraindications": [],
            "fatigue_score":     {},
            # raw source
            "source_data": ex["details"],
        })

    total = len(exercises)
    print(f"\nProcessing {total} exercises in batches of {BATCH_SIZE} …\n")

    # 5. LLM enrichment in batches
    for start in range(0, total, BATCH_SIZE):
        batch_ex = exercises[start:start + BATCH_SIZE]
        pct = (start + len(batch_ex)) / total * 100
        names_str = ", ".join(e["name"] for e in batch_ex)
        print(f"[{start+len(batch_ex):3d}/{total}] ({pct:5.1f}%)  {names_str[:80]}")

        enrichments = call_llm(llm, batch_ex)

        for ex, enr in zip(batch_ex, enrichments):
            ex["name_cn"]           = enr.get("name_cn")           or ex["name"]
            ex["secondary_muscles"] = safe_list(enr.get("secondary_muscles"))
            ex["common_mistakes"]   = safe_list(enr.get("common_mistakes"))
            ex["safety_tips"]       = safe_list(enr.get("safety_tips"))
            ex["contraindications"] = safe_list(enr.get("contraindications"))
            ex["estimated_mets"]    = safe_float(enr.get("estimated_mets"), ex["estimated_mets"])
            ex["fatigue_score"]     = safe_dict(enr.get("fatigue_score"))

        time.sleep(0.3)  # gentle rate limit

    # 6. Upsert into PostgreSQL
    print("\nInserting into PostgreSQL …")
    ok = err = 0
    with conn.cursor() as cur:
        for ex in exercises:
            try:
                params = {k: json.dumps(v, ensure_ascii=False)
                          if isinstance(v, (dict, list)) else v
                          for k, v in ex.items()}
                cur.execute(UPSERT_SQL, params)
                ok += 1
            except Exception as e:
                print(f"  [DB] Error for '{ex['name']}': {e}")
                conn.rollback()
                err += 1
                continue
        conn.commit()

    conn.close()
    print(f"\nDone! Inserted/updated: {ok}  Errors: {err}")

    # 7. Quick sanity check
    conn2 = psycopg2.connect(**DB_CFG)
    with conn2.cursor() as cur:
        cur.execute("SELECT COUNT(*) FROM exercises")
        cnt = cur.fetchone()[0]
        cur.execute("SELECT exercise_id, name, name_cn, primary_muscles FROM exercises LIMIT 5")
        rows = cur.fetchall()
    conn2.close()
    print(f"\nTotal rows in exercises table: {cnt}")
    print("Sample rows:")
    for row in rows:
        print(f"  {row[0]}: {row[1]} / {row[2]} | muscles={row[3]}")


if __name__ == "__main__":
    main()
