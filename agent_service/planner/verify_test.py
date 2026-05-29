# -*- coding: utf-8 -*-
import openpyxl, json

wb = openpyxl.load_workbook("D:/WorkPath/fitness/agent_service/planner/workout_details_test.xlsx")
ws = wb.active
print("Rows:", ws.max_row, " Cols:", ws.max_column)
print("Headers:", [ws.cell(1, c).value for c in range(1, 5)])
print()

for r in range(2, ws.max_row + 1):
    plan  = ws.cell(r, 1).value
    name  = ws.cell(r, 2).value
    steps = (ws.cell(r, 3).value or "")[:120].replace("\n", " ")
    det   = ws.cell(r, 4).value or "{}"
    d     = json.loads(det)
    bi    = d.get("basic_info", {})
    mf    = d.get("muscles", {}).get("frontBodyMap", {}).get("text-mw-red", [])
    mb    = d.get("muscles", {}).get("backBodyMap",  {}).get("text-mw-red", [])
    print(f"  [{r-1:2d}] Plan : {plan}")
    print(f"       Name : {name}")
    print(f"       Steps: {steps}...")
    print(f"       Info : {bi}")
    print(f"       Muscles front-red: {mf}")
    print(f"       Muscles back-red : {mb}")
    print()
