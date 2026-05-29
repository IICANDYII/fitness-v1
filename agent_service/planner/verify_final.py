# -*- coding: utf-8 -*-
"""Verify and summarize workout_details.xlsx"""
import openpyxl, json, collections

wb = openpyxl.load_workbook("D:/WorkPath/fitness/agent_service/planner/workout_details.xlsx")
ws = wb.active
print(f"Rows (incl header): {ws.max_row}")
print(f"Cols: {ws.max_column}")
print(f"Headers: {[ws.cell(1,c).value for c in range(1,5)]}")
print()

# Count exercises per plan
plan_counts = collections.Counter()
for r in range(2, ws.max_row + 1):
    plan_counts[ws.cell(r, 1).value] += 1

print(f"Total plans: {len(plan_counts)}")
print(f"Total exercises: {sum(plan_counts.values())}")
print()
print("Sample rows (first 3, last 3):")
for r in [2, 3, 4, ws.max_row-2, ws.max_row-1, ws.max_row]:
    plan  = ws.cell(r, 1).value
    name  = ws.cell(r, 2).value
    steps = (ws.cell(r, 3).value or "")[:80].replace("\n", " ")
    det   = ws.cell(r, 4).value or "{}"
    d     = json.loads(det)
    bi    = d.get("basic_info", {})
    mf    = d.get("muscles", {}).get("frontBodyMap", {}).get("text-mw-red", [])
    mb    = d.get("muscles", {}).get("backBodyMap",  {}).get("text-mw-red", [])
    print(f"  Row {r}: {plan} | {name}")
    print(f"    Steps: {steps}...")
    print(f"    Info:  {bi}")
    print(f"    Muscles(red): front={mf}  back={mb}")
    print()

# Show any rows with empty steps
empty = [r for r in range(2, ws.max_row+1) if not (ws.cell(r,3).value or "").strip()]
if empty:
    print(f"Rows with empty steps: {len(empty)}")
    for r in empty[:5]:
        print(f"  Row {r}: {ws.cell(r,1).value} | {ws.cell(r,2).value}")
