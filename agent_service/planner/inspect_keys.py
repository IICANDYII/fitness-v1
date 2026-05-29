# -*- coding: utf-8 -*-
"""Inspect the actual JSON key bytes in basic_info to handle encoding correctly."""
import openpyxl, json, sys

wb = openpyxl.load_workbook("D:/WorkPath/fitness/agent_service/planner/workout_details.xlsx")
ws = wb.active

# Collect all unique key-value combos
all_info = {}
for r in range(2, min(ws.max_row + 1, 100)):
    det = ws.cell(r, 4).value or "{}"
    d = json.loads(det)
    bi = d.get("basic_info", {})
    for k, v in bi.items():
        if v not in ["", None]:
            all_info.setdefault(k, set()).add(v)

print("Keys and their value sets:")
for k, vs in all_info.items():
    k_repr = repr(k)
    print(f"  {k_repr}  ->  {sorted(vs)[:8]}")
